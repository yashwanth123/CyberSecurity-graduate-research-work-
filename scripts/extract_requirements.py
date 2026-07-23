#!/usr/bin/env python3
"""Extract IEC 62443 compliance data from Excel workbook to JSON."""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "_IEC 62443__.xlsx"
OUT = ROOT / "webapp" / "src" / "data" / "iec62443.json"


def extract() -> dict:
    wb = openpyxl.load_workbook(XLSX, read_only=True)

    catalog = []
    for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:]:
        if r[0]:
            catalog.append({
                "part": r[0],
                "abbreviation": r[1],
                "standard": r[2],
                "role": r[3],
                "description": (r[4] or "") if len(r) > 4 else "",
            })

    questions = []
    for i, r in enumerate(list(wb["Sheet2"].iter_rows(values_only=True))[1:], start=1):
        r = list(r) + [None] * 10
        if r[4]:
            questions.append({
                "id": f"Q{i:03d}",
                "part": r[0],
                "subcategory": r[1],
                "standard": r[2],
                "framework_category": r[3],
                "question": str(r[4]).strip(),
                "sl1": r[6], "sl2": r[7], "sl3": r[8], "sl4": r[9],
            })

    sl_matrix = []
    for r in list(wb["Sheet3"].iter_rows(values_only=True))[1:]:
        if r[0] is not None:
            sl_matrix.append({
                "level": int(r[0]),
                "misuse": r[1], "means": r[2], "resources": r[3],
                "knowledge": r[4], "motivation": r[5],
            })

    return {"catalog": catalog, "questions": questions, "slMatrix": sl_matrix}


def main():
    if not XLSX.exists():
        print(f"Excel file not found: {XLSX}")
        sys.exit(1)
    data = extract()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, indent=2))
    print(f"Wrote {len(data['questions'])} questions to {OUT}")


if __name__ == "__main__":
    main()
